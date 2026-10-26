# -*- coding: utf-8 -*-
"""Build a synthetic "school regulation" DOCX used as a template-channel fixture.

The sample deliberately differs from our own layout: 仿宋 12pt body, 黑体 table
headings, a footer, a 5-column stage table with one data row plus a 备注 footer
row, and label/value cells in fixed positions. Tests assert that filling this
file keeps every one of those traits.
"""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BODY_FONT = "仿宋"
HEAD_FONT = "黑体"


def _run(paragraph, text, font, size, bold=False):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    return run


def _line(document, text="", font=BODY_FONT, size=12, bold=False, align=None):
    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(3)
    _run(paragraph, text, font, size, bold)
    return paragraph


def _table(document, rows, cols, widths=None):
    table = document.add_table(rows=rows, cols=cols)
    table.style = "Table Grid"
    table.alignment = 1
    if widths:
        table.autofit = False
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    return table


def _fill(cell, text, font=BODY_FONT, size=12, bold=False):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    _run(paragraph, text, font, size, bold)
    return cell


def build_school_template(path, data_rows=1):
    """Write the fixture DOCX and return its path.

    ``data_rows`` controls how many blank stage rows the sample offers; a sample
    with more blank rows than the lesson needs is a real-world case worth testing.
    """
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2.2)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2.2)

    _line(document, "××学院教案", HEAD_FONT, 22, True, WD_ALIGN_PARAGRAPH.CENTER)
    _line(document, "")

    info = _table(document, 5, 4, [2.6, 5.5, 2.7, 5.5])
    labels = [("课程名称", "课程学分"), ("授课课题", "授课班级"),
              ("授课教师", "授课周次"), ("学时", "课型"),
              ("授课日期", "教学地点")]
    for row_index, (left, right) in enumerate(labels):
        _fill(info.cell(row_index, 0), left, HEAD_FONT, 10.5, True)
        _fill(info.cell(row_index, 1), "")
        _fill(info.cell(row_index, 2), right, HEAD_FONT, 10.5, True)
        _fill(info.cell(row_index, 3), "")

    _line(document, "")
    _line(document, "一、教学目标", HEAD_FONT, 14, True)
    objectives = _table(document, 3, 2, [3.2, 13.1])
    for row_index, label in enumerate(["知识目标", "能力目标", "素质目标"]):
        _fill(objectives.cell(row_index, 0), label, HEAD_FONT, 10.5, True)
        _fill(objectives.cell(row_index, 1), "")

    _line(document, "")
    _line(document, "二、学情分析", HEAD_FONT, 14, True)
    _line(document, "", BODY_FONT, 12)
    _line(document, "")
    _line(document, "三、教学重点与难点", HEAD_FONT, 14, True)
    points = _table(document, 2, 2, [3.2, 13.1])
    for row_index, label in enumerate(["教学重点", "教学难点"]):
        _fill(points.cell(row_index, 0), label, HEAD_FONT, 10.5, True)
        _fill(points.cell(row_index, 1), "")

    _line(document, "")
    _line(document, "四、教学过程", HEAD_FONT, 14, True)
    stage = _table(document, 2 + data_rows, 5, [2.5, 5.3, 3.1, 3.0, 2.4])
    for col_index, header in enumerate(["教学环节", "教学内容", "学生活动", "教师活动", "设计意图"]):
        _fill(stage.cell(0, col_index), header, HEAD_FONT, 10.5, True)
    for row_index in range(1, 1 + data_rows):
        for col_index in range(5):
            _fill(stage.cell(row_index, col_index), "")
    footer_row = stage.rows[1 + data_rows]
    merged = footer_row.cells[0].merge(footer_row.cells[4])
    _fill(merged, "备注：本表行数可根据实际教学环节增减。", BODY_FONT, 9)

    _line(document, "")
    _line(document, "五、教学反思", HEAD_FONT, 14, True)
    reflection = _table(document, 3, 2, [3.2, 13.1])
    for row_index, label in enumerate(["教学效果", "教学诊断", "教学改进"]):
        _fill(reflection.cell(row_index, 0), label, HEAD_FONT, 10.5, True)
        _fill(reflection.cell(row_index, 1), "")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(footer, "××学院教务处制", BODY_FONT, 9)

    header = section.header.paragraphs[0]
    _run(header, "教案（样本）", HEAD_FONT, 9)

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    document.save(path)
    return path


def build_unit_style_template(path, units=7):
    """Fixture for the other common layout: one small table under each stage heading."""
    document = Document()
    section = document.sections[0]
    section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.2)
    section.left_margin, section.right_margin = Cm(2.0), Cm(2.0)

    _line(document, "第 X 天 教案", HEAD_FONT, 16, True, WD_ALIGN_PARAGRAPH.CENTER)
    _line(document, "一、基本信息", HEAD_FONT, 14, True)
    info = _table(document, 2, 2, [3.0, 13.3])
    _fill(info.cell(0, 0), "课程名称", HEAD_FONT, 10.5, True)
    _fill(info.cell(0, 1), "")
    _fill(info.cell(1, 0), "授课课题", HEAD_FONT, 10.5, True)
    _fill(info.cell(1, 1), "")

    _line(document, "二、教学目标", HEAD_FONT, 14, True)
    for index in range(1, 4):
        _line(document, "%d. 样本目标占位第%d条" % (index, index))

    _line(document, "三、教学过程", HEAD_FONT, 14, True)
    for index in range(units):
        _line(document, "环节·单元%d（知识点：样本知识点%d）" % (index + 1, index + 1), HEAD_FONT, 11.5, True)
        table = _table(document, 2, 4, [4.0, 3.8, 3.8, 4.0])
        for column, header in enumerate(["教学内容", "学生活动", "教师活动", "设计意图"]):
            _fill(table.cell(0, column), header, HEAD_FONT, 10, True)
            _fill(table.cell(1, column), "")

    _line(document, "四、教学反思", HEAD_FONT, 14, True)
    reflection = _table(document, 3, 2, [3.0, 13.3])
    for row, label in enumerate(["教学效果", "诊断", "改进"]):
        _fill(reflection.cell(row, 0), label, HEAD_FONT, 10.5, True)
        _fill(reflection.cell(row, 1), "")

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    document.save(path)
    return path


if __name__ == "__main__":
    import sys

    target = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "fixtures", "school_template.docx")
    print(build_school_template(target))
