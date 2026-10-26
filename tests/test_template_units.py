# -*- coding: utf-8 -*-
"""Units mode: documents that put one small table under each stage heading."""
import os
import shutil
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, "scripts")
STACK = os.path.join(ROOT, "examples", "stack_queue", "lesson.json")
sys.path.insert(0, SCRIPTS)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document  # noqa: E402
from common import load_lesson  # noqa: E402
from docx_template import (  # noqa: E402
    analyze_sample, numbered_ranges, render_from_template,
)
from template_fixture import BODY_FONT, build_unit_style_template  # noqa: E402


class UnitsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.data = load_lesson(STACK)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def sample(self, units=7):
        path = os.path.join(self.tmp, "units_%d.docx" % units)
        return build_unit_style_template(path, units=units)

    def render(self, units, data=None):
        sample = self.sample(units)
        mapping, _ = analyze_sample(sample)
        out = os.path.join(self.tmp, "out_%d.docx" % units)
        summary = render_from_template(sample, data or self.data, mapping, out)
        return mapping, summary, Document(out)

    def test_analysis_detects_units_mode(self):
        mapping, report = analyze_sample(self.sample(7))
        block = mapping["blocks"]["in_class"]
        self.assertEqual(block["mode"], "units")
        self.assertEqual(block["unit_count"], 7)
        self.assertEqual(block["columns"]["content"], 0)
        self.assertIn("知识点", block["heading_template"])
        self.assertTrue(any(x["kind"] == "block" and x.get("mode") == "units"
                            for x in report["candidates"]))

    def test_units_are_paired_when_counts_match(self):
        mapping, summary, document = self.render(7)
        self.assertEqual(summary["units"], 7)
        self.assertEqual(summary["cloned_units"], 0)
        self.assertEqual(summary["removed_units"], 0)
        headings = [p.text for p in document.paragraphs if p.text.startswith("环节·")]
        self.assertEqual(len(headings), 7)
        self.assertIn("导入", headings[1])
        self.assertIn("课前", headings[0])
        self.assertIn("课后", headings[-1])
        self.assertEqual(len(document.tables), 9)   # 基本信息 + 7 单元 + 反思

    def test_units_are_cloned_when_sample_has_fewer(self):
        mapping, summary, document = self.render(3)
        self.assertEqual(summary["cloned_units"], 4)
        self.assertEqual(len([p for p in document.paragraphs if p.text.startswith("环节·")]), 7)
        stages = [s["stage"] for s in self.data["flow"]["in_class"]]
        headings = "\n".join(p.text for p in document.paragraphs if p.text.startswith("环节·"))
        for stage in stages:
            self.assertIn(stage, headings)
        cloned_table = document.tables[4]
        self.assertTrue(cloned_table.rows[1].cells[0].text.strip())

    def test_surplus_units_are_removed(self):
        mapping, summary, document = self.render(9)
        self.assertEqual(summary["removed_units"], 2)
        self.assertEqual(len([p for p in document.paragraphs if p.text.startswith("环节·")]), 7)

    def test_unit_cells_keep_sample_font(self):
        mapping, summary, document = self.render(7)
        unit_table = document.tables[1]
        run = unit_table.rows[1].cells[0].paragraphs[0].runs[0]
        self.assertEqual(run.font.name, BODY_FONT)
        self.assertEqual(run.font.size.pt, 12)
        self.assertTrue(document.tables[1].rows[1].cells[0].text.strip())

    def test_numbered_objectives_become_range(self):
        sample = self.sample(7)
        document = Document(sample)
        ranges = numbered_ranges(document)
        self.assertEqual(ranges[0], (3, 5))          # 三条编号目标合并成一个区间
        mapping, _ = analyze_sample(sample)
        target = mapping["fields"]["objectives.all"]
        self.assertEqual((target["kind"], target["start"], target["end"]),
                         ("paragraph_range", 3, 5))
        out = os.path.join(self.tmp, "objectives.docx")
        render_from_template(sample, self.data, mapping, out)
        text = "\n".join(p.text for p in Document(out).paragraphs)
        self.assertIn("知识目标", text)
        self.assertIn("能力目标", text)
        self.assertNotIn("样本目标占位", text)


if __name__ == "__main__":
    unittest.main()
