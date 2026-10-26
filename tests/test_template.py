# -*- coding: utf-8 -*-
"""Template channel: sample analysis, format-preserving fill, and failure paths."""
import copy
import json
import os
import shutil
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, "scripts")
STACK = os.path.join(ROOT, "examples", "stack_queue", "lesson.json")
sys.path.insert(0, SCRIPTS)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document  # noqa: E402
from common import LessonInputError, load_lesson  # noqa: E402
from docx_template import (  # noqa: E402
    analyze_sample, field_value, load_map, render_from_template,
)
from template_fixture import BODY_FONT, HEAD_FONT, build_school_template  # noqa: E402


def cell_run(cell):
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            if run.text.strip():
                return run
    return None


class TemplateCase(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.sample = os.path.join(self.tmp, "school_sample.docx")
        build_school_template(self.sample)
        self.data = load_lesson(STACK)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def render(self, mapping, data=None):
        out = os.path.join(self.tmp, "filled.docx")
        summary = render_from_template(self.sample, data or self.data, mapping, out)
        return out, summary, Document(out)


class AnalysisTests(TemplateCase):
    def test_detects_meta_fields_and_block(self):
        mapping, report = analyze_sample(self.sample)
        fields = mapping["fields"]
        for key in ["lesson_meta.course_name", "lesson_meta.teacher", "lesson_meta.week",
                    "lesson_meta.hours", "objectives.knowledge", "objectives.ability",
                    "objectives.quality", "key_points", "difficult_points", "student_analysis"]:
            self.assertIn(key, fields, key)
        self.assertIn("in_class", mapping["blocks"])
        block = mapping["blocks"]["in_class"]
        self.assertEqual(block["template_row"], 1)
        self.assertEqual(block["last_row"], 1)          # 备注行被识别为表尾，不参与填充
        self.assertEqual(block["columns"]["content"], 1)
        self.assertEqual(block["columns"]["design_intent"], 4)
        self.assertEqual(block["columns"]["stage_with_duration"], 0)
        self.assertTrue(report["candidates"])

    def test_report_lists_what_could_not_be_located(self):
        mapping, report = analyze_sample(self.sample)
        self.assertIsInstance(report["ambiguous"], list)
        self.assertNotIn("lesson_meta.course_name", report["ambiguous"])

    def test_analyze_rejects_missing_file(self):
        with self.assertRaises(LessonInputError):
            analyze_sample(os.path.join(self.tmp, "nope.docx"))


class FillTests(TemplateCase):
    def mapping(self):
        mapping, _ = analyze_sample(self.sample)
        return mapping

    def test_fills_meta_and_keeps_sample_fonts(self):
        out, summary, document = self.render(self.mapping())
        info = document.tables[0]
        value_cell = info.cell(0, 1)
        self.assertEqual(value_cell.text.strip(), "数据结构")
        run = cell_run(value_cell)
        self.assertIsNotNone(run)
        self.assertEqual(run.font.name, BODY_FONT)       # 仿宋 12 来自样本，未被替换
        self.assertEqual(run.font.size.pt, 12)
        label_run = cell_run(info.cell(0, 0))
        self.assertEqual(label_run.font.name, HEAD_FONT)
        self.assertEqual(summary["fields"], len(self.mapping()["fields"]))

    def test_expands_stage_rows_and_keeps_footer_row(self):
        out, summary, document = self.render(self.mapping())
        stage = document.tables[3]
        self.assertEqual(summary["cloned_rows"], 4)      # 1 行扩到 5 个环节
        self.assertEqual(len(stage.rows), 7)             # 表头 + 5 环节 + 备注
        stages = [s["stage"] for s in self.data["flow"]["in_class"]]
        for index, name in enumerate(stages, 1):
            self.assertIn(name, stage.cell(index, 0).text, name)
        self.assertIn("备注", stage.cell(6, 0).text)
        cloned_run = cell_run(stage.cell(3, 1))
        self.assertEqual(cloned_run.font.name, BODY_FONT)
        self.assertEqual(cloned_run.font.size.pt, 12)

    def test_objectives_and_body_paragraph(self):
        out, _, document = self.render(self.mapping())
        knowledge = document.tables[1].cell(0, 1).text
        self.assertTrue(knowledge.startswith("1. "))
        self.assertIn("\n", knowledge)
        analysis = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("学生已学习线性表", analysis)

    def test_does_not_touch_styles_header_or_footer(self):
        out, _, document = self.render(self.mapping())
        self.assertEqual(document.sections[0].footer.paragraphs[0].text, "××学院教务处制")
        self.assertEqual(document.sections[0].header.paragraphs[0].text, "教案（样本）")
        with zipfile.ZipFile(self.sample) as a, zipfile.ZipFile(out) as b:
            self.assertEqual(a.read("word/styles.xml"), b.read("word/styles.xml"))
        margins = document.sections[0]
        self.assertAlmostEqual(margins.top_margin.cm, 2.5, places=2)

    def test_map_roundtrip_via_json(self):
        map_path = os.path.join(self.tmp, "map.json")
        mapping = self.mapping()
        mapping["confirmed"] = True
        with open(map_path, "w", encoding="utf-8") as f:
            json.dump(mapping, f, ensure_ascii=False)
        loaded = load_map(map_path)
        out, _, document = self.render(loaded)
        self.assertEqual(document.tables[0].cell(1, 3).text.strip(), "计科2401班")

    def test_more_items_than_sample_rows_clones_exactly(self):
        mapping = self.mapping()
        out, summary, document = self.render(mapping)
        self.assertEqual(summary["removed_rows"], 0)

    def test_block_trims_rows_when_data_is_shorter(self):
        # 学校样本给了 3 行空位，本节课只有 1 个环节：多余空行必须删掉
        roomy = os.path.join(self.tmp, "school_sample_3rows.docx")
        build_school_template(roomy, data_rows=3)
        mapping, _ = analyze_sample(roomy)
        self.assertEqual(mapping["blocks"]["in_class"]["last_row"], 3)
        shorter = copy.deepcopy(self.data)
        shorter["flow"]["in_class"] = shorter["flow"]["in_class"][:1]
        out = os.path.join(self.tmp, "trimmed.docx")
        summary = render_from_template(roomy, shorter, mapping, out)
        document = Document(out)
        self.assertEqual(summary["removed_rows"], 2)
        self.assertEqual(len(document.tables[3].rows), 3)  # 表头 + 1 环节 + 备注
        self.assertIn("备注", document.tables[3].cell(2, 0).text)

    def test_repeat_header_and_no_row_split_options(self):
        from docx.oxml.ns import qn
        mapping = self.mapping()
        mapping["blocks"]["in_class"]["repeat_header"] = True
        mapping["blocks"]["in_class"]["keep_row_together"] = True
        out, _, document = self.render(mapping)
        stage = document.tables[3]
        header_trpr = stage.rows[0]._tr.find(qn("w:trPr"))
        self.assertIsNotNone(header_trpr)
        self.assertIsNotNone(header_trpr.find(qn("w:tblHeader")))
        for row in stage.rows[1:6]:
            self.assertIsNotNone(row._tr.find(qn("w:trPr")).find(qn("w:cantSplit")))

    def test_options_absent_keeps_sample_defaults(self):
        from docx.oxml.ns import qn
        out, _, document = self.render(self.mapping())
        stage = document.tables[3]
        header_trpr = stage.rows[0]._tr.find(qn("w:trPr"))
        self.assertTrue(header_trpr is None or header_trpr.find(qn("w:tblHeader")) is None)

    def test_bad_address_fails_clearly(self):
        mapping = self.mapping()
        mapping["fields"]["lesson_meta.course_name"] = {"kind": "cell", "table": 0, "row": 99, "col": 1}
        with self.assertRaises(LessonInputError) as ctx:
            self.render(mapping)
        self.assertIn("不存在", str(ctx.exception))

    def test_unsupported_column_key_fails_clearly(self):
        mapping = self.mapping()
        mapping["blocks"]["in_class"]["columns"]["bogus"] = 2
        with self.assertRaises(LessonInputError) as ctx:
            self.render(mapping)
        self.assertIn("bogus", str(ctx.exception))

    def test_empty_mapping_is_rejected(self):
        with self.assertRaises(LessonInputError):
            self.render({"fields": {}, "blocks": {}})

    def test_field_value_helpers(self):
        self.assertTrue(field_value("objectives.knowledge", self.data).startswith("1. "))
        self.assertIn("栈", field_value("in_class_summary", self.data))
        self.assertEqual(field_value("flow.total_minutes", self.data), "90")
        self.assertTrue(field_value("knowledge_outline", self.data))
        self.assertEqual(field_value("lesson_meta.week", self.data), "第15周")


if __name__ == "__main__":
    unittest.main()
