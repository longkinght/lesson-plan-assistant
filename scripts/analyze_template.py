# -*- coding: utf-8 -*-
"""Analyze a school-provided DOCX and propose a reusable template map."""
import argparse
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import LessonInputError, load_lesson  # noqa: E402
from docx_template import (  # noqa: E402
    analyze_sample, field_value, load_map, map_report_markdown, resolve_address,
)
from docx import Document  # noqa: E402


def main():
    parser = argparse.ArgumentParser(description="解析学校/教师 Word 样本，生成可复用的模板映射")
    parser.add_argument("--template", required=True, help="样本 .docx 路径")
    parser.add_argument("--out-dir", help="映射与报告输出目录，默认与样本同目录")
    parser.add_argument("--json", help="可选：用一份 lesson.json 预检字段是否有内容")
    parser.add_argument("--check", action="store_true", help="只校验已有映射是否仍与样本匹配")
    parser.add_argument("--map", help="配合 --check 使用：已有 template-map.json")
    args = parser.parse_args()

    try:
        if args.check:
            if not args.map:
                raise LessonInputError("--check 需要同时指定 --map")
            mapping = load_map(args.map)
            document = Document(args.template)
            problems = []
            for key, address in (mapping.get("fields") or {}).items():
                resolve_address(document, address, key)
            for source, spec in (mapping.get("blocks") or {}).items():
                table_index = int(spec.get("table", -1))
                if table_index < 0 or table_index >= len(document.tables):
                    problems.append("blocks.%s 指向的表不存在" % source)
                    continue
                table = document.tables[table_index]
                if int(spec.get("template_row", 0)) >= len(table.rows):
                    problems.append("blocks.%s 的模板行超出表格范围" % source)
            if problems:
                print("映射与样本不匹配：")
                for item in problems:
                    print("  - %s" % item)
                return 1
            print("映射与样本匹配：%d 个字段、%d 个区块"
                  % (len(mapping.get("fields") or {}), len(mapping.get("blocks") or {})))
            return 0

        mapping, report = analyze_sample(args.template)
        out_dir = os.path.abspath(args.out_dir or os.path.dirname(os.path.abspath(args.template)))
        os.makedirs(out_dir, exist_ok=True)
        stem = os.path.splitext(os.path.basename(args.template))[0]
        map_path = os.path.join(out_dir, stem + "_template-map.json")
        report_path = os.path.join(out_dir, stem + "_template-report.md")
        with open(map_path, "w", encoding="utf-8") as f:
            json.dump(mapping, f, ensure_ascii=False, indent=2)
        with open(report_path, "w", encoding="utf-8") as f:
            f.write(map_report_markdown(mapping, report))

        print("样本解析完成：%s" % report["template"])
        print("  段落 %d · 表格 %d · 节 %d"
              % (report["paragraphs"], report["tables"], report["sections"]))
        print("  识别字段 %d 个，重复区块 %d 个"
              % (len(mapping["fields"]), len(mapping["blocks"])))
        if report["ambiguous"]:
            print("  未定位（需手工补充）：%s" % "、".join(report["ambiguous"]))
        print("  映射文件：%s" % map_path)
        print("  核对报告：%s" % report_path)

        if args.json:
            data = load_lesson(args.json)
            empty = [key for key in mapping["fields"] if not field_value(key, data)]
            if empty:
                print("  预检：以下字段在 lesson.json 中无内容，将写入空白")
                for key in empty:
                    print("    - %s" % key)
            else:
                print("  预检：全部字段都能从 lesson.json 取到内容")
        print("首次使用前请打开映射文件核对，并把 confirmed 改为 true。")
        return 0
    except (LessonInputError, ValueError, OSError) as exc:
        print("解析失败：%s" % exc)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
