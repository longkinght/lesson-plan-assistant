# -*- coding: utf-8 -*-
"""Audit the three rendered artifacts against their single source of truth."""
import argparse
import json
import os
import re
import sys

from docx import Document

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (  # noqa: E402
    LessonInputError, load_lesson, objective_items, objective_text,
    safe_file_stem, validate_lesson,
)


def extract_html_data(html):
    match = re.search(r'<script id="lesson-data" type="application/json">(.*?)</script>', html, re.S)
    if not match:
        raise ValueError("HTML 中缺少 lesson-data 数据块")
    return json.loads(match.group(1))


def audit(json_path, output_dir):
    data = load_lesson(json_path)
    errors, warnings = validate_lesson(data)
    stem = safe_file_stem(data["lesson_meta"])
    paths = {
        "markdown": os.path.join(output_dir, stem + "_知识图谱.md"),
        "html": os.path.join(output_dir, stem + "_知识图谱.html"),
        "docx": os.path.join(output_dir, stem + "_教案.docx"),
    }
    checks = []
    source_nodes = [node["name"] for node in data["knowledge_tree"] if node["node_type"] == "知识点"]

    for kind, path in paths.items():
        checks.append({"name": "%s 文件存在" % kind, "ok": os.path.isfile(path), "detail": path})

    if os.path.isfile(paths["markdown"]):
        with open(paths["markdown"], "r", encoding="utf-8") as f:
            markdown = f.read()
        missing = [name for name in source_nodes if name not in markdown]
        checks.append({"name": "Markdown 覆盖全部知识点", "ok": not missing, "detail": ", ".join(missing) or "全部命中"})

    if os.path.isfile(paths["html"]):
        with open(paths["html"], "r", encoding="utf-8") as f:
            html = f.read()
        try:
            embedded = extract_html_data(html)
            html_nodes = [n["name"] for n in embedded.get("knowledge_tree", []) if n.get("node_type") == "知识点"]
            checks.append({"name": "HTML 内嵌知识与源数据一致", "ok": html_nodes == source_nodes,
                           "detail": "%d / %d" % (len(html_nodes), len(source_nodes))})
        except (ValueError, json.JSONDecodeError) as exc:
            checks.append({"name": "HTML 内嵌知识与源数据一致", "ok": False, "detail": str(exc)})

    if os.path.isfile(paths["docx"]):
        document = Document(paths["docx"])
        text = "\n".join(p.text for p in document.paragraphs)
        for table in document.tables:
            text += "\n" + "\n".join(cell.text for row in table.rows for cell in row.cells)
        missing_stages = [stage["stage"] for stage in data["flow"]["in_class"] if stage["stage"] not in text]
        checks.append({"name": "DOCX 包含全部课中环节", "ok": not missing_stages, "detail": ", ".join(missing_stages) or "全部命中"})
        top_objectives = [objective_text(item) for group in ["knowledge", "ability", "quality"]
                          for item in objective_items(data, group)]
        missing_goals = [goal for goal in top_objectives if goal and goal not in text]
        checks.append({"name": "DOCX 包含三维教学目标", "ok": not missing_goals,
                       "detail": "%d 个缺失" % len(missing_goals)})

    report = {"ok": not errors and all(item["ok"] for item in checks),
              "errors": errors, "warnings": warnings, "checks": checks}
    print("=== v2 产物验收 ===")
    for item in checks:
        print("[%s] %s：%s" % ("通过" if item["ok"] else "失败", item["name"], item["detail"]))
    if errors:
        print("契约错误：")
        for item in errors:
            print("  - %s" % item)
    return report


def main():
    parser = argparse.ArgumentParser(description="核查 DOCX、Markdown、HTML 与 lesson.json 的同源一致性")
    parser.add_argument("--json", required=True, help="lesson.json 路径")
    parser.add_argument("--output-dir", required=True, help="三件套输出目录")
    parser.add_argument("--report", help="可选 JSON 报告路径")
    args = parser.parse_args()
    try:
        report = audit(args.json, args.output_dir)
    except (LessonInputError, ValueError, OSError) as exc:
        print("验收失败：%s" % exc)
        return 2
    if args.report:
        with open(args.report, "w", encoding="utf-8") as f:
            json.dump(report, f, ensure_ascii=False, indent=2)
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
