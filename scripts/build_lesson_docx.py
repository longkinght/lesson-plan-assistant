# -*- coding: utf-8 -*-
"""Build a polished, editable DOCX lesson plan from the validated lesson JSON."""
import argparse
import json
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common import (
    LessonInputError, build_tree, generate_post_relations, load_lesson,
    objective_items, objective_refs, objective_text, validate_lesson,
)
import docx_template

HEAD_FONT = "微软雅黑"
BODY_FONT = "宋体"
INK = RGBColor(0x22, 0x32, 0x38)
TEAL = RGBColor(0x0E, 0x5A, 0x5E)
TEAL_LIGHT = "E8F2EF"
GOLD = RGBColor(0xA5, 0x78, 0x2E)
GOLD_LIGHT = "F7EFDD"
MUTED = RGBColor(0x6E, 0x80, 0x83)
LINE = "D8E2DD"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
STYLE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                         "references", "teacher-style.json")


def load_style(style, config_path=None):
    if config_path:
        try:
            with open(config_path, "r", encoding="utf-8-sig") as f:
                payload = json.load(f)
        except (OSError, json.JSONDecodeError) as exc:
            raise LessonInputError("无法读取风格配置：%s" % exc)
        if not isinstance(payload, dict):
            raise LessonInputError("风格配置顶层必须是对象")
        return payload
    if style in ["chapter", "teacher-zhang"]:
        try:
            with open(STYLE_DIR, "r", encoding="utf-8-sig") as f:
                return json.load(f)
        except (OSError, json.JSONDecodeError) as exc:
            raise LessonInputError("无法读取内置章节风格：%s" % exc)
    if style != "standard":
        raise LessonInputError("未知风格：%s（可选 standard、chapter，或使用 --style-config）" % style)
    return None


def set_run(run, font=BODY_FONT, size=10.5, bold=False, color=None, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    return run


def add_para(doc, text="", font=BODY_FONT, size=10.5, bold=False, color=None,
             align=None, before=0, after=4, line=1.35, keep=False, italic=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    if text:
        set_run(p.add_run(str(text)), font, size, bold, color, italic)
    return p


def add_bottom_border(paragraph, color=LINE, size="8", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_heading(doc, index, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(index + "  "), HEAD_FONT, 9, True, GOLD)
    set_run(p.add_run(text), HEAD_FONT, 13.5, True, TEAL)
    add_bottom_border(p, LINE)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), size)
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), color)
        borders.append(e)
    tbl_pr.append(borders)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_text(cell, text, font=BODY_FONT, size=9.5, bold=False, color=INK, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.32
    set_run(p.add_run("" if text is None else str(text)), font, size, bold, color)
    set_cell_margins(cell)


def add_kv_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = (Cm(3.4), Cm(13.6))
    for key, value in rows:
        cells = table.add_row().cells
        keep_row_together(table.rows[-1])
        for index, cell in enumerate(cells):
            cell.width = widths[index]
        set_cell_text(cells[0], key, HEAD_FONT, 9.5, True, TEAL)
        set_cell_shading(cells[0], TEAL_LIGHT)
        set_cell_text(cells[1], value, BODY_FONT, 9.5, False, INK)
    return table


def add_info_grid(doc, meta):
    values = [
        ("课程名称", meta.get("course_name", "")), ("授课课题", meta.get("title", "")),
        ("周次", meta.get("week", "")), ("学时", meta.get("hours", "")),
        ("授课班级", meta.get("class_name", "")), ("授课教师", meta.get("teacher", "")),
        ("授课日期", meta.get("date", "")), ("课型", meta.get("course_type", "")),
    ]
    table = doc.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [Cm(2.5), Cm(5.9), Cm(2.5), Cm(6.1)]
    for index, (label, value) in enumerate(values):
        row_i, pair_i = divmod(index, 2)
        col = pair_i * 2
        label_cell, value_cell = table.cell(row_i, col), table.cell(row_i, col + 1)
        label_cell.width, value_cell.width = widths[col], widths[col + 1]
        set_cell_text(label_cell, label, HEAD_FONT, 8.5, True, MUTED)
        set_cell_shading(label_cell, "F5F8F6")
        set_cell_text(value_cell, value or "—", BODY_FONT, 9.2)
    for row in table.rows:
        keep_row_together(row)
    return table


def add_stage_table(doc, stage, index=None, duration=None):
    title = stage.get("stage", "")
    refs = stage.get("knowledge_refs", [])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    if index is not None:
        set_run(p.add_run("%02d  " % index), HEAD_FONT, 9, True, GOLD)
    set_run(p.add_run(title), HEAD_FONT, 10.5, True, TEAL)
    if duration:
        set_run(p.add_run("  /  %s 分钟" % duration), HEAD_FONT, 8.5, False, MUTED)
    if refs:
        set_run(p.add_run("　%s" % " · ".join(refs)), BODY_FONT, 8, False, MUTED)

    headers = ["教学内容", "学生活动", "教师活动", "设计意图"]
    dimensions = ["content", "student_activity", "teacher_activity", "design_intent"]
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = (Cm(4.9), Cm(3.7), Cm(3.7), Cm(4.7))
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = widths[i]
        set_cell_text(cell, header, HEAD_FONT, 8.5, True, WHITE, WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(cell, "0E5A5E")
    set_repeat_header(table.rows[0])
    for i, dimension in enumerate(dimensions):
        cell = table.rows[1].cells[i]
        cell.width = widths[i]
        set_cell_text(cell, stage.get(dimension, ""), BODY_FONT, 9.2, False, INK)
    keep_row_together(table.rows[1])
    add_para(doc, "", after=1)


def add_objectives(doc, data):
    for group, label in [("knowledge", "知识目标"), ("ability", "能力目标"), ("quality", "素质目标")]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        set_run(p.add_run(label), HEAD_FONT, 9.5, True, TEAL)
        for number, item in enumerate(objective_items(data, group), 1):
            text = objective_text(item)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.36
            set_run(p.add_run("%d. " % number), BODY_FONT, 10, True, TEAL)
            set_run(p.add_run(text), BODY_FONT, 10, False, INK)
            refs = objective_refs(item)
            if refs:
                set_run(p.add_run("　[%s]" % " / ".join(refs)), BODY_FONT, 8, False, MUTED)


def add_page_fields(paragraph):
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)
    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    instr_run._r.append(instr)
    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def configure_document(doc, data):
    meta = data["lesson_meta"]
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(1.9)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
    section.header_distance, section.footer_distance = Cm(0.9), Cm(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.paragraph_format.space_after = Pt(0)
    set_run(header.add_run("LESSON PLAN"), HEAD_FONT, 8, True, TEAL)
    set_run(header.add_run("  /  %s" % meta.get("course_name", "")), HEAD_FONT, 8, False, MUTED)
    add_bottom_border(header, LINE, "4", "2")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("AI 辅助生成 · 教师审核后使用　　"), BODY_FONT, 7.5, False, MUTED)
    add_page_fields(footer)

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    props = doc.core_properties
    props.title = "%s %s 教案" % (meta.get("week", ""), meta.get("title", ""))
    props.subject = meta.get("course_name", "")
    props.author = meta.get("teacher", "") or "Lesson Plan Assistant"
    props.comments = "由 lesson.json 同源生成；提交前请完成教学适用性与事实审核。"
    props.keywords = "教案, 知识图谱, 课程设计"


def render_cover(doc, meta):
    add_para(doc, "COURSE LESSON PLAN", HEAD_FONT, 8.5, True, GOLD, after=10)
    add_para(doc, meta.get("title", "教案"), HEAD_FONT, 25, True, TEAL, after=3, line=1.12)
    subtitle = "  ·  ".join([x for x in [meta.get("course_name", ""), meta.get("week", ""), meta.get("course_type", "")] if x])
    add_para(doc, subtitle, BODY_FONT, 10, False, MUTED, after=8)
    line = add_para(doc, "", after=10)
    add_bottom_border(line, "B08A3E", "14", "1")
    add_info_grid(doc, meta)


def render_standard(doc, data, meta):
    render_cover(doc, meta)

    add_heading(doc, "01", "教学目标")
    add_objectives(doc, data)

    add_heading(doc, "02", "学情分析与教学重难点")
    add_para(doc, "学情分析", HEAD_FONT, 9.5, True, TEAL, before=3, after=3, keep=True)
    add_para(doc, data.get("student_analysis") or "待补充学情分析。", after=7)
    add_kv_table(doc, [
        ("教学重点", "；".join(data.get("key_points", [])) or "待补充"),
        ("教学难点", "；".join(data.get("difficult_points", [])) or "待补充"),
    ])

    add_heading(doc, "03", "教学方法与资源")
    add_kv_table(doc, [
        ("教学方法", "；".join(data.get("methods", [])) or "待补充"),
        ("教学资源", "；".join(data.get("resources", [])) or "待补充"),
    ])

    add_heading(doc, "04", "教学过程")
    flow = data["flow"]
    add_para(doc, "课前 · 自主学习", HEAD_FONT, 10, True, TEAL, before=2, after=3, keep=True)
    add_stage_table(doc, flow["pre_class"])
    add_para(doc, "课中 · 教学实施", HEAD_FONT, 10, True, TEAL, before=7, after=3, keep=True)
    for index, stage in enumerate(flow["in_class"], 1):
        add_stage_table(doc, stage, index, stage.get("duration_minutes"))
    add_para(doc, "课后 · 巩固拓展", HEAD_FONT, 10, True, TEAL, before=7, after=3, keep=True)
    add_stage_table(doc, flow["post_class"])

    add_heading(doc, "05", "教学反思")
    reflection = data["reflection"]
    status = reflection.get("status", "pre_class")
    note = "本课尚未实施，以下为课后记录占位。" if status == "pre_class" else "以下内容应来自课堂记录与作业结果。"
    add_para(doc, note, BODY_FONT, 8.5, False, MUTED, after=5, italic=True)
    add_kv_table(doc, [
        ("教学效果", reflection.get("effect", "")),
        ("诊断", reflection.get("diagnosis", "")),
        ("改进", reflection.get("improvement", "")),
    ])

    add_heading(doc, "附录", "知识结构与目标索引")
    roots, _ = build_tree(data["knowledge_tree"])
    add_para(doc, "本页由 lesson.json 自动生成，用于核对教案所依据的知识结构；正式提交时可按学校要求删除。",
             BODY_FONT, 8.5, False, MUTED, after=6, italic=True)
    for root in roots:
        render_outline(doc, root)


def render_outline(doc, node, depth=0):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(depth * 0.55)
    p.paragraph_format.space_before = Pt(2 if depth == 0 else 0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.25
    is_category = node.get("node_type") == "分类"
    set_run(p.add_run(node.get("name", "")), HEAD_FONT if is_category else BODY_FONT, 9 if depth < 2 else 8.5,
            is_category, TEAL if is_category else INK)
    if not is_category:
        details = [x for x in [node.get("cognitive", ""), node.get("category", "")] if x]
        if details:
            set_run(p.add_run("　%s" % " / ".join(details)), BODY_FONT, 7.5, False, MUTED)
    for child in node.get("children", []):
        render_outline(doc, child, depth + 1)


def render_chapter(doc, data, style, meta):
    cover = style.get("cover", {})
    chapter = style.get("chapter", {})
    defaults = chapter.get("defaults", {})
    add_para(doc, "CHAPTER LESSON PLAN", HEAD_FONT, 8.5, True, GOLD, after=10)
    title = str(cover.get("title_template", "{course_name}教案")).replace("{course_name}", meta.get("course_name", "课程"))
    add_para(doc, title, HEAD_FONT, 24, True, TEAL, after=3, line=1.12)
    semester = cover.get("semester") or meta.get("semester", "")
    add_para(doc, semester, BODY_FONT, 10, False, MUTED, after=9)
    line = add_para(doc, "", after=10)
    add_bottom_border(line, "B08A3E", "14", "1")
    cover_rows = []
    for field in cover.get("fields", []):
        template = str(field.get("value", ""))
        match = re.fullmatch(r"\{(\w+)\}", template)
        value = meta.get(match.group(1), "") if match else template
        cover_rows.append((field.get("label", ""), value or field.get("fallback", "") or "—"))
    if cover_rows:
        add_kv_table(doc, cover_rows)
    if cover.get("note"):
        add_para(doc, cover["note"], BODY_FONT, 8, False, MUTED, after=8, italic=True)

    chapter_num = meta.get("chapter_num") or (re.search(r"\d+", meta.get("week", "")).group(0) if re.search(r"\d+", meta.get("week", "")) else "1")
    heading = str(chapter.get("head_template", "第{chapter_num}章")).replace("{chapter_num}", str(chapter_num))
    add_heading(doc, "01", "%s  %s" % (heading, meta.get("title", "")))

    def field_block(label, content, size=9.5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        set_run(p.add_run(label), HEAD_FONT, 9.5, True, TEAL)
        values = content if isinstance(content, list) else [content]
        for value in values:
            add_para(doc, value, BODY_FONT, size, after=2, line=1.4)

    for field in chapter.get("field_order", []):
        if field == "课时":
            field_block(field, "%s 学时" % meta.get("hours", ""))
        elif field == "教学方法":
            methods = data.get("methods", [])
            joined = "".join(methods)
            field_block(field, "　".join(("■" if any(k in joined for k in option.split("/")) else "□") + option
                                         for option in ["讲授法", "讨论法", "案例教学", "参与式教学", "其他"]))
        elif field == "教学目的及要求":
            knowledge = [objective_text(x) for x in objective_items(data, "knowledge")]
            ability = [objective_text(x) for x in objective_items(data, "ability")]
            quality = [objective_text(x) for x in objective_items(data, "quality")]
            field_block(field, ["目的：" + "；".join(knowledge + ability) + "。",
                                "要求：" + "；".join(quality) + "。"])
        elif field == "教学环境":
            field_block(field, (data.get("resources") or [defaults.get("教学环境", "多媒体教室")])[0])
        elif field == "重点":
            field_block(field, "重点：" + "；".join(data.get("key_points", [])))
        elif field == "难点":
            field_block(field, "难点：" + "；".join(data.get("difficult_points", [])))
        elif field == "讲授内容":
            lines = []
            for index, stage in enumerate(data["flow"]["in_class"], 1):
                lines.append("%d. %s（%s 分钟）" % (index, stage.get("stage", ""), stage.get("duration_minutes", "")))
                lines.append(stage.get("content", ""))
                if stage.get("student_activity"):
                    lines.append("学生活动：" + stage["student_activity"])
            field_block(field, lines)
        elif field == "思考题和作业":
            post = data["flow"]["post_class"]
            field_block(field, [post.get("content", ""), "学生活动：" + post.get("student_activity", "")])
        elif field == "参考文献":
            field_block(field, data.get("references") or defaults.get("参考文献", ""))

    add_heading(doc, "02", "教学反思（附加）")
    reflection = data["reflection"]
    add_para(doc, "本课尚未实施，以下为课后记录占位。" if reflection.get("status") == "pre_class"
             else "以下内容应来自课堂记录与作业结果。", BODY_FONT, 8.5, False, MUTED, after=5, italic=True)
    add_kv_table(doc, [("教学效果", reflection.get("effect", "")),
                       ("诊断", reflection.get("diagnosis", "")),
                       ("改进", reflection.get("improvement", ""))])


def main_inner(json_path, output_path, style="standard", style_config=None,
               template=None, template_map=None):
    data = load_lesson(json_path)
    errors, _ = validate_lesson(data)
    if errors:
        print("校验未通过，拒绝生成 DOCX：")
        for item in errors:
            print("  - %s" % item)
        return 1
    generate_post_relations(data["knowledge_tree"])

    if template:
        mapping = None
        if template_map:
            mapping = docx_template.load_map(template_map)
        else:
            mapping, report = docx_template.analyze_sample(template)
            auto_path = os.path.join(os.path.dirname(os.path.abspath(output_path)),
                                     os.path.splitext(os.path.basename(template))[0] + "_template-map.json")
            os.makedirs(os.path.dirname(auto_path), exist_ok=True)
            with open(auto_path, "w", encoding="utf-8") as f:
                json.dump(mapping, f, ensure_ascii=False, indent=2)
            print("未提供 --template-map，已自动推断映射：%s" % auto_path)
            if report.get("ambiguous"):
                print("  ⚠ 以下字段未能定位，需人工补充：%s" % "、".join(report["ambiguous"]))
            print("  ⚠ 自动推断仅供首次试用；正式交付前必须核对映射并置 confirmed=true。")
        summary = docx_template.render_from_template(template, data, mapping, output_path)
        detail = "字段 %d · 区块 %d · 新增行 %d · 删除行 %d" % (
            summary["fields"], summary["blocks"], summary["cloned_rows"], summary["removed_rows"])
        if summary.get("units"):
            detail += " · 单元 %d（克隆 %d / 删除 %d）" % (
                summary["units"], summary["cloned_units"], summary["removed_units"])
        print("DOCX 已生成：%s（严格样本格式：%s · %s）" % (output_path, summary["template"], detail))
        if not mapping.get("confirmed"):
            print("  提示：该映射尚未标记 confirmed，请核对 %s"
                  % (template_map or "同目录自动生成的 template-map.json"))
        return 0

    style_data = load_style(style, style_config)
    meta = data["lesson_meta"]
    doc = Document()
    configure_document(doc, data)
    if style_data:
        render_chapter(doc, data, style_data, meta)
        label = "章节模板"
    else:
        render_standard(doc, data, meta)
        label = "标准教案"
    doc.save(output_path)
    print("DOCX 已生成：%s（%s）" % (output_path, label))
    return 0


def main():
    parser = argparse.ArgumentParser(description="教案 DOCX 生成器（standard / chapter / 样本模板）")
    parser.add_argument("--json", required=True, help="lesson.json 路径")
    parser.add_argument("--output", required=True, help="输出 .docx 路径")
    parser.add_argument("--style", default="standard", help="standard、chapter 或旧版别名 teacher-zhang")
    parser.add_argument("--style-config", help="自定义章节风格 JSON")
    parser.add_argument("--template", help="学校/教师提供的样本 .docx，按其实有格式严格填充")
    parser.add_argument("--template-map", help="已确认的 template-map.json；缺省则自动推断")
    args = parser.parse_args()
    try:
        return main_inner(args.json, args.output, args.style, args.style_config,
                          args.template, args.template_map)
    except (LessonInputError, ValueError, OSError) as exc:
        print("生成失败：%s" % exc)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
